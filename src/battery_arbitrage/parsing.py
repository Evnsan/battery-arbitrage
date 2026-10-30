import math
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from battery_arbitrage.domain import MarketData

PRICE_HEADER_SUFFIX = "Price [£/MWh]"
MONTH_FIRST_TIMESTAMP_FORMAT = "%m/%d/%y %H:%M"


def parse_market_data(file_path: str | Path) -> MarketData:
    """
    Parse market price data from an Excel workbook (.xlsx).

    The workbook contains one worksheet per market. The worksheet name is used
    as the market name. In each worksheet, timestamps are read from the first
    column, which does not need a header. The price column is identified by a
    header ending with "Price [£/MWh]". Timestamps may be Excel datetimes,
    ISO-formatted strings, or month-first strings in "MM/DD/YY HH:MM" format.
    Fully blank rows are ignored.

    This function only converts the file into MarketData.

    Args:
        file_path (str | Path): Path to the Excel workbook.

    Returns:
        MarketData: A dictionary mapping each market name (worksheet name) to a
            dictionary of interval timestamps and prices in £/MWh.

    Raises:
        FileNotFoundError: If file_path does not exist.
        ValueError: If the file cannot be read as an Excel workbook, or if a
            worksheet has no data rows, lacks a price column, has more than one
            price column, contains a missing or unparseable timestamp, contains
            a missing, unparseable, or non-finite price, or contains duplicate
            timestamps. The message names the offending worksheet.
    """
    try:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
    except FileNotFoundError:
        raise
    except Exception as error:
        raise ValueError(
            f"Could not read Excel workbook: {file_path}"
        ) from error

    market_data: MarketData = {}
    try:
        for worksheet in workbook.worksheets:
            sheet_name = worksheet.title
            rows = worksheet.iter_rows(values_only=True)
            header_row = None
            header_row_number = 0

            for row_number, row in enumerate(rows, start=1):
                if any(value is not None for value in row):
                    header_row = row
                    header_row_number = row_number
                    break

            if header_row is None:
                raise ValueError(f"Worksheet '{sheet_name}' has no data rows")

            price_columns = [
                column_index
                for column_index, heading in enumerate(header_row)
                if isinstance(heading, str)
                and heading.endswith(PRICE_HEADER_SUFFIX)
            ]
            if not price_columns:
                raise ValueError(
                    f"Worksheet '{sheet_name}' "
                    f"is missing a price column ending in "
                    f"'{PRICE_HEADER_SUFFIX}'"
                )
            if len(price_columns) > 1:
                raise ValueError(
                    f"Worksheet '{sheet_name}' has more than one price column"
                )

            price_column = price_columns[0]
            market_prices: dict[datetime, float] = {}
            for row_number, row in enumerate(rows, start=header_row_number + 1):
                if not any(value is not None for value in row):
                    continue

                timestamp_value = row[0] if row else None
                price_value = (
                    row[price_column] if price_column < len(row) else None
                )
                if timestamp_value is None:
                    raise ValueError(
                        f"Worksheet '{sheet_name}' has a missing timestamp "
                        f"on row {row_number}"
                    )
                if price_value is None:
                    raise ValueError(
                        f"Worksheet '{sheet_name}' has a missing price "
                        f"on row {row_number}"
                    )

                if isinstance(timestamp_value, datetime):
                    timestamp = timestamp_value
                elif isinstance(timestamp_value, str):
                    try:
                        timestamp = datetime.fromisoformat(timestamp_value)
                    except ValueError:
                        try:
                            timestamp = datetime.strptime(
                                timestamp_value, MONTH_FIRST_TIMESTAMP_FORMAT
                            )
                        except ValueError as error:
                            raise ValueError(
                                f"Worksheet '{sheet_name}' "
                                f"has an unparseable timestamp "
                                f"on row {row_number}: {timestamp_value!r}"
                            ) from error
                else:
                    raise ValueError(
                        f"Worksheet '{sheet_name}' "
                        f"has an unparseable timestamp "
                        f"on row {row_number}: {timestamp_value!r}"
                    )

                if isinstance(price_value, bool):
                    raise ValueError(
                        f"Worksheet '{sheet_name}' has an unparseable price "
                        f"on row {row_number}: {price_value!r}"
                    )
                if not isinstance(price_value, (int, float, Decimal, str)):
                    raise ValueError(
                        f"Worksheet '{sheet_name}' has an unparseable price "
                        f"on row {row_number}: {price_value!r}"
                    )
                try:
                    price = float(price_value)
                except (TypeError, ValueError) as error:
                    raise ValueError(
                        f"Worksheet '{sheet_name}' has an unparseable price "
                        f"on row {row_number}: {price_value!r}"
                    ) from error
                if not math.isfinite(price):
                    raise ValueError(
                        f"Worksheet '{sheet_name}' has a non-finite price "
                        f"on row {row_number}: {price_value!r}"
                    )

                if timestamp in market_prices:
                    raise ValueError(
                        f"Worksheet '{sheet_name}' has duplicate timestamp "
                        f"{timestamp!r}"
                    )
                market_prices[timestamp] = price

            if not market_prices:
                raise ValueError(f"Worksheet '{sheet_name}' has no data rows")

            market_data[sheet_name] = market_prices
    finally:
        workbook.close()

    return market_data
